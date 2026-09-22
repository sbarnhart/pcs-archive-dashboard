import collections
import openpyxl

report_path = r"C:\Users\steve\OneDrive\Desktop\StockMasterList (2).xlsx"
import_path = r"outputs\embed-product-import-20260922\Inventory product template - Embed populated.xlsx"

def txt(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

report = openpyxl.load_workbook(report_path, read_only=True, data_only=True).active
report.reset_dimensions()
report_rows = list(report.iter_rows(values_only=True))[4:-1]
import_sheet = openpyxl.load_workbook(import_path, read_only=True, data_only=True).active
import_rows = list(import_sheet.iter_rows(values_only=True))[1:]
by_barcode = collections.defaultdict(list)
for row in import_rows:
    if txt(row[10]):
        by_barcode[txt(row[10])].append(row)
report_matches = []
unmatched = []
multi = []
for row in report_rows:
    codes = [x.strip() for x in txt(row[8]).split(";") if x.strip()]
    if len(codes) > 1:
        multi.append((row[1], codes))
    hits = [imp for code in codes for imp in by_barcode.get(code, [])]
    if hits:
        report_matches.append((row, hits))
    else:
        unmatched.append((row[1], row[8]))
print("report_rows", len(report_rows))
print("report_barcodes_multi", len(multi))
print("report_matched_rows", len(report_matches))
print("report_unmatched_rows", len(unmatched))
print("nonzero_sale_price", sum(1 for row in report_rows if isinstance(row[5], (int, float)) and row[5] != 0))
print("nonblank_uom", sum(1 for row in report_rows if txt(row[7])))
print("type_counts", collections.Counter(txt(row[3]) for row in report_rows))
print("sale_price_examples", [(r[1],r[5],r[8]) for r in report_rows if isinstance(r[5], (int,float)) and r[5] != 0][:15])
print("multi_examples", multi[:10])
print("unmatched_examples", unmatched[:15])
print("matched_name_mismatch", [(r[1],hits[0][1],r[8]) for r,hits in report_matches if txt(r[1]).casefold()!=txt(hits[0][1]).casefold()][:15])
